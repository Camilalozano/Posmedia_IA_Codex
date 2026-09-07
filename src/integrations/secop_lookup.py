"""Consulta por referencia contractual en una exportación Excel o CSV.

El archivo se procesa en memoria. La conexión Oracle utilizará este mismo
esquema cuando se restablezca el acceso y se verifiquen sus columnas.
"""
import csv
import hashlib
import io
import re
import unicodedata
import zipfile
import socket
from dataclasses import dataclass
from datetime import date, datetime
from urllib.error import HTTPError, URLError
from urllib.parse import urlparse
from urllib.request import Request, urlopen

from openpyxl import load_workbook
from src.config import NOT_FOUND
from src.models import Field

REFERENCE = 'referencia_del_contrato (contratos_electronicos)'
ADVANCE = 'porc_ejecucion_financiera'
SHEET = 'tabla_maestra_completa'
FIELD_MAP = {
    'numero_contrato_convenio': REFERENCE,
    'nombre_contratista_asociado': 'proveedor_adjudicado (contratos_electronicos)',
    'nombre_supervisor': 'nombre_supervisor (contratos_electronicos)',
    'fecha_terminacion': 'fecha_de_fin_del_contrato (contratos_electronicos)',
    'modificaciones': 'tipo_modificacion',
    'objeto': 'descripcion_del_proceso (contratos_electronicos)',
}
MAX_BYTES = 50 * 1024 * 1024
MAX_EXPANDED_BYTES = 250 * 1024 * 1024
MAX_ROWS = 200_000
ORACLE_HOST = 'objectstorage.us-ashburn-1.oraclecloud.com'


class OracleConnectionError(RuntimeError):
    """Error seguro para mostrar sin revelar el PAR configurado."""


def oracle_connection_message(detail=''):
    suffix = f' ({detail})' if detail else ''
    return (
        'No fue posible conectar con la base de Oracle' + suffix + '. '
        'Revisa e ingresa un nuevo PAR en la configuración Secrets de Streamlit, '
        'con el nombre ORACLE_PAR_URL, y vuelve a intentar.'
    )


@dataclass
class LookupResult:
    reference: str
    fields: dict[str, Field]
    warnings: list[str]
    source: str
    row_number: int
    sha256: str
    raw_advance: str = ''


def normalize_reference(value):
    """Ignora caja, espacios y guiones; conserva dígitos y demás caracteres.

No elimina ceros iniciales ni realiza coincidencias parciales o aproximadas.
"""
    if value is None:
        return ''
    text = unicodedata.normalize('NFKC', str(value)).upper().strip()
    return re.sub(r'[\s\-\u2010-\u2015\u2212]+', '', text)


def cell_text(value):
    if value is None:
        return ''
    if isinstance(value, (date, datetime)):
        return value.isoformat()
    text = str(value).strip()
    if text.casefold() in {'nan', 'nat', 'null', 'none', 'no definido', 'no identificado'}:
        return ''
    return text


def validate_headers(headers):
    names = [cell_text(value) for value in headers]
    if REFERENCE not in names:
        raise ValueError('La base no contiene la columna requerida: ' + REFERENCE)
    nonempty = [name for name in names if name]
    if len(nonempty) != len(set(nonempty)):
        raise ValueError('La base contiene encabezados duplicados; revise el archivo.')
    return names


def match_rows(headers, rows, query):
    key = normalize_reference(query)
    if not key:
        raise ValueError('Indique el número del proceso antes de consultar.')
    names = validate_headers(headers)
    position = names.index(REFERENCE)
    found = None
    for row_number, row in enumerate(rows, 2):
        if row_number > MAX_ROWS + 1:
            raise ValueError('La base supera el límite de 200.000 registros.')
        if position >= len(row) or normalize_reference(row[position]) != key:
            continue
        if found is not None:
            raise ValueError('Hay varios registros con la misma referencia normalizada. Revise los duplicados de la base; no se ha seleccionado ninguno.')
        found = (row_number, dict(zip(names, row)))
    return found


def map_record(row, source, row_number, digest):
    reference = cell_text(row[REFERENCE])
    fields, warnings = {}, []
    for target, column in FIELD_MAP.items():
        raw = cell_text(row.get(column))
        provenance = f'{source} · fila {row_number} · {column}'
        if not raw:
            fields[target] = Field(NOT_FOUND, provenance + ' · sin dato', 'bajo')
            warnings.append(f'La base no informa: {target.replace("_", " ")}.')
            continue
        value = raw
        if target == 'fecha_terminacion':
            try:
                value = datetime.fromisoformat(raw).strftime('%d/%m/%Y')
            except ValueError:
                warnings.append('Revise el formato de la fecha de terminación; se conservó el valor original.')
        fields[target] = Field(value, provenance, 'dato de base; por verificar')
    raw_advance = cell_text(row.get(ADVANCE))
    provenance = f'{source} · fila {row_number} · {ADVANCE}'
    if raw_advance:
        fields['porcentaje_avance'] = Field(raw_advance, provenance, 'dato de base; por verificar')
    else:
        fields['porcentaje_avance'] = Field(NOT_FOUND, provenance + ' · sin dato', 'bajo')
        warnings.append('La base no informa el porcentaje de ejecución financiera.')
    return LookupResult(reference, fields, warnings, source, row_number, digest, raw_advance)


def lookup_file(filename, data, query):
    """Devuelve un registro único, None si no existe o un error explícito."""
    if len(data) > MAX_BYTES:
        raise ValueError('La base supera el límite de 50 MB.')
    digest = hashlib.sha256(data).hexdigest()
    suffix = filename.rsplit('.', 1)[-1].lower()
    if suffix == 'xlsx':
        try:
            with zipfile.ZipFile(io.BytesIO(data)) as archive:
                if sum(item.file_size for item in archive.infolist()) > MAX_EXPANDED_BYTES:
                    raise ValueError('El Excel descomprimido supera 250 MB.')
            workbook = load_workbook(io.BytesIO(data), read_only=True, data_only=True, keep_links=False)
        except (zipfile.BadZipFile, OSError):
            raise ValueError('No se pudo leer el Excel. Revise que sea un archivo .xlsx válido.') from None
        try:
            if SHEET not in workbook.sheetnames:
                raise ValueError('El Excel debe contener la hoja ' + SHEET + '.')
            sheet = workbook[SHEET]
            if sheet.max_column and sheet.max_column > 500:
                raise ValueError('La hoja supera 500 columnas.')
            iterator = sheet.iter_rows(values_only=True)
            found = match_rows(next(iterator, []), iterator, query)
        finally:
            workbook.close()
        source = filename + ' · ' + SHEET
    elif suffix == 'csv':
        try:
            text = data.decode('utf-8-sig')
        except UnicodeDecodeError:
            raise ValueError('El CSV debe estar codificado en UTF-8.') from None
        try:
            dialect = csv.Sniffer().sniff(text[:32768], delimiters=',;\t')
        except csv.Error:
            dialect = csv.excel
        iterator = csv.reader(io.StringIO(text, newline=''), dialect=dialect)
        found = match_rows(next(iterator, []), iterator, query)
        source = filename
    else:
        raise ValueError('Cargue una base Excel (.xlsx) o CSV.')
    if found is None:
        return None
    row_number, row = found
    return map_record(row, source, row_number, digest)


def download_oracle_csv(par_url, timeout=30, opener=urlopen):
    """Descarga el CSV configurado, con límites y mensajes que no exponen el PAR."""
    parsed = urlparse(str(par_url).strip())
    if parsed.scheme != 'https' or parsed.hostname != ORACLE_HOST:
        raise OracleConnectionError(oracle_connection_message('PAR inválido'))
    request = Request(str(par_url).strip(), headers={'User-Agent': 'Posmedia-IA-Codex/0.3'})
    try:
        with opener(request, timeout=timeout) as response:
            declared = response.headers.get('Content-Length')
            if declared:
                try:
                    if int(declared) > MAX_BYTES:
                        raise OracleConnectionError(oracle_connection_message('el archivo supera 50 MB'))
                except ValueError:
                    raise OracleConnectionError(oracle_connection_message('respuesta inválida')) from None
            data = response.read(MAX_BYTES + 1)
            if len(data) > MAX_BYTES:
                raise OracleConnectionError(oracle_connection_message('el archivo supera 50 MB'))
    except OracleConnectionError:
        raise
    except HTTPError as error:
        raise OracleConnectionError(oracle_connection_message('código HTTP ' + str(error.code))) from None
    except (URLError, TimeoutError, socket.timeout, OSError):
        raise OracleConnectionError(oracle_connection_message('error de red o tiempo de espera')) from None
    if not data:
        raise OracleConnectionError(oracle_connection_message('la respuesta está vacía'))
    return data


def lookup_oracle(par_url, query, opener=urlopen):
    data = download_oracle_csv(par_url, opener=opener)
    try:
        return lookup_file('Oracle Object Storage · tabla_maestra_completa.csv', data, query)
    except (UnicodeDecodeError, ValueError) as error:
        # El enlace puede apuntar a un objeto distinto o a una exportación incompatible.
        raise OracleConnectionError(oracle_connection_message('archivo incompatible: ' + str(error))) from None
